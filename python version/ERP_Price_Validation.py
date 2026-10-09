# %%
import re
from copy import copy
import datetime
import pandas as pd
import openpyxl as opxl
from openpyxl.utils import column_index_from_string

dfNFs = pd.read_excel("PowerQuery DMV.xlsx", sheet_name="Query1")
dfEPR = pd.read_excel("PowerQuery ERP.xlsx", sheet_name="NetPrices")
dfEX = pd.read_excel("PowerQuery Exchange Rates.xlsx", sheet_name="ExRates")
dfRefs = pd.read_excel("PowerQuery Referencial.xlsx", sheet_name="Referencial")

# %%
df = dfNFs[1:]

# %%
pns = df['Detail by N.F. Number']
descriptions = df['Column2']
idco = df['Column3']
supid = df['Column16']
supname = df['Column17']
volume = df['Column18']
currency = df['Column20']
price = df['Column27']

compare = []

for i in range(len(pns)):
    val_supid = str(supid.iloc[i]).split('.')[0].strip() if pd.notna(supid.iloc[i]) else ""
    compare.append(str(pns.iloc[i]) + "|" + val_supid)

# %%
pnsEPR = dfEPR['PN EPR']
supidEPR = dfEPR['ID EPR']
moedaEPR = dfEPR['MOEDA EPR']
priceEPR = dfEPR['PRICE EPR']
umEPR = dfEPR['UM EPR']

compareEPR: dict[str,int] = {}
for i in range(len(pnsEPR)):
    val_supidEPR = str(supidEPR.iloc[i]).split('.')[0].strip() if pd.notna(supidEPR.iloc[i]) else ""
    compareEPR[str(pnsEPR.iloc[i]) + "|" + val_supidEPR] = i

# %%
currencyEX = dfEX['CurrencyCode']
ratesEX = dfEX['Rates2026']

compareEX: dict[str,float] = {}
for i in range(len(currencyEX)):
    compareEX[str(currencyEX.iloc[i])] = ratesEX.iloc[i]

# %%
add = []

# %%
import datetime
print(datetime.datetime.now().date())

# %%
print(len(dfRefs['PN']))
print(str(dfRefs['ID'].iloc[1390]))

# %%
dictRG = {}
dictCM = {}
for i in range(len(dfRefs)):
    dictRG[str(dfRefs['PN'].iloc[i])] = str(dfRefs['REGIONAL / GLOBAL'].iloc[i])
    if pd.notna(dfRefs['ID'].iloc[i]) and str(dfRefs['ID'].iloc[i]) != "nan":
        dictCM[str(dfRefs['ID'].iloc[i]).split('.')[0].strip()] = str(dfRefs['CM'].iloc[i])

# %%
print(dictCM)

# %%
for i in range(len(compare)):
    if compare[i] in compareEPR:
        indexEPR = compareEPR[compare[i]]

        if currency.iloc[i] == 'BRL':
            loc = 'L'
        elif str(supname.iloc[i])[0:5].lower() == 'husqv':
            loc = 'I'
        else:
            loc = 'O'

        if umEPR.iloc[indexEPR] == 'MI':
            pricePC = priceEPR.iloc[indexEPR]/1000
        elif umEPR.iloc[indexEPR] == 'CT':
            pricePC = priceEPR.iloc[indexEPR]/100
        else:
            pricePC = priceEPR.iloc[indexEPR]
        
        rate_moedaEPR = compareEX.get(str(moedaEPR.iloc[indexEPR]), 1.0)
        rate_currency = compareEX.get(str(currency.iloc[i]), 1.0)

        priceEPRBRL = priceEPR.iloc[indexEPR]*rate_moedaEPR
        priceBRL = price.iloc[i]*rate_currency
        
        dif_pct = abs((priceBRL - priceEPRBRL)/priceBRL) if priceBRL != 0 else 0.0
        val_supid_lookup = str(supid.iloc[i]).split('.')[0].strip() if pd.notna(supid.iloc[i]) else ""

        add.append([str(pns.iloc[i]),str(pns.iloc[i])[0:9],price.iloc[i],currency.iloc[i],priceBRL,"",datetime.datetime.now().date(),volume.iloc[i],price.iloc[i]*volume.iloc[i],priceBRL*volume.iloc[i],supid.iloc[i],supname.iloc[i],dictCM.get(val_supid_lookup, 'Not Found'),abs(priceBRL- priceEPRBRL),dif_pct,pnsEPR.iloc[indexEPR],priceEPR.iloc[indexEPR],pricePC,moedaEPR.iloc[indexEPR],priceEPRBRL,priceEPR.iloc[indexEPR]*volume.iloc[i],priceEPRBRL*volume.iloc[i],umEPR.iloc[indexEPR],supidEPR.iloc[indexEPR],descriptions.iloc[i],"",dictRG.get(str(pns.iloc[i]), 'Not Found'),loc,idco.iloc[i]])

# %%
def add_df_data_to_existing_excel_table(excel_file, sheet_name, table_name, df_new_data):
    wb = opxl.load_workbook(excel_file)
    ws = wb[sheet_name]
    table = ws.tables[table_name]

    start_cell, end_cell = table.ref.split(":")
    start_col_letter = re.match(r"([A-Za-z]+)", start_cell).group(1)
    start_col_idx = column_index_from_string(start_col_letter) 

    current_table_end_row = int(re.search(r"\d+", end_cell).group())
    row_num = current_table_end_row + 1
    
    column_key = {name: index + start_col_idx for index, name in enumerate(table.column_names)}

    for d in df_new_data.to_dict(orient='records'):
        for col, value in d.items():
            if col in column_key:
                col_idx = column_key[col]
                cell = ws.cell(row=row_num, column=col_idx)
                cell.value = value

                prev_cell = ws.cell(row=row_num - 1, column=col_idx)
                if prev_cell.has_style:
                    cell.font = copy(prev_cell.font)
                    cell.border = copy(prev_cell.border)
                    cell.fill = copy(prev_cell.fill)
                    cell.number_format = copy(prev_cell.number_format)
                    cell.alignment = copy(prev_cell.alignment)

        row_num += 1

    end_col_letters = re.match(r"([A-Za-z]+)", end_cell).group(1)
    table.ref = f"{start_cell}:{end_col_letters}{row_num - 1}"

    wb.save(excel_file)

# %%
excel_file = "priceDivergence.xlsx"
sheet_name = "Main"
table_name = "PriceDivergenceTable"
df_new_data = pd.DataFrame(add, columns=['PN - CONTROLLER', 'PN S/ C', 'PRICE - CONTROLLER', 'CURRENCY - CONTROLLER', 'ACTUAL PRICE IN BRL', 'MEASURE - ACTUAL', 'DATA FATURAMENTO', 'VOLUME', 'SPEND - ACTUAL', 'SPEND - ACTUAL BRL', 'SUPPLIER ID CONTROLLER', 'SUPPLIER NAME', 'CM', '|DIFERENÇA BRL|', '|DIFERENÇA%|', 'PN - EPR', 'PRICE - EPR', 'PRICE - EPR PC', 'CURRENCY - EPR', 'EPR PRICE IN BRL', 'SPEND - EPR', 'SPEND - EPR BRL', 'MEASURE - EPR', 'SUPPLIER EPR', 'DESCRIPTION', 'COMPARE SUP&PRICE', 'R / G', 'L /O / I', 'IDCO'])

add_df_data_to_existing_excel_table(excel_file, sheet_name, table_name, df_new_data)